import asyncio
import io
from datetime import datetime, timedelta
from decimal import Decimal

import pytz
from aiogram import Bot
from aiogram.types import BufferedInputFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import settings
from database.connection import async_session_factory
from database.crud import get_export_tickets

HEADERS = [
    "ID", "Дата", "Время", "Телефон", "Дополнительный контакт", "Модель",
    "Источник", "Статус клиента", "Мастер", "Время сеанса", "Длительность (мин)",
    "Результат", "Комментарий",
    "Оплата за отмену", "Оплата сеанса", "Консумация",
    "Доп. время (мин)", "Оплата за доп. время", "Оператор",
]

CLIENT_STATUS_MAP = {
    "new": "Новый",
    "white_list": "Белый список",
    "black_list": "Чёрный список",
}

RESULT_MAP = {
    "positive_lead": "Позитивный лид",
    "booking": "Бронь",
    "general_lead": "Общий лид",
}


def _ticket_to_row(ticket, tz) -> list:
    """Convert a Ticket ORM object to a plain list. Must be called while session is open."""
    pay: dict = {"session": Decimal(0), "cancellation": Decimal(0),
                 "service": Decimal(0), "extra_time": Decimal(0)}
    extra_mins = 0
    for p in ticket.payments:
        ptype = p.payment_type.value
        pay[ptype] = pay.get(ptype, Decimal(0)) + p.amount
        if ptype == "extra_time" and p.extra_time_minutes:
            extra_mins += p.extra_time_minutes

    created = ticket.created_at
    if created.tzinfo is None:
        created = pytz.utc.localize(created)
    created = created.astimezone(tz)

    comment = ticket.cancellation_reason or ticket.departure_reason or ""
    client_status = CLIENT_STATUS_MAP.get(
        ticket.client_status.value if ticket.client_status else "new", "Новый"
    )
    result = RESULT_MAP.get(
        ticket.result_category.value if ticket.result_category else "", ""
    )

    def _pay(key):
        v = float(pay[key])
        return v if v else ""

    return [
        ticket.id,
        created.strftime("%d.%m.%Y"),
        created.strftime("%H:%M"),
        ticket.client_phone or "",
        ticket.client_contact or "",
        ticket.model.name if ticket.model else "",
        ticket.traffic_source.name if ticket.traffic_source else "",
        client_status,
        ticket.master.name if ticket.master else "",
        ticket.session_start or "",
        ticket.session_duration or "",
        result,
        comment,
        _pay("cancellation"),
        _pay("session"),
        _pay("service"),
        extra_mins or "",
        _pay("extra_time"),
        ticket.operator.full_name if ticket.operator else "",
    ]


def _generate_excel(rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Выгрузка"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E4057")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    ws.row_dimensions[1].height = 36

    alt_fill = PatternFill("solid", fgColor="EEF2F7")
    for row_idx, row in enumerate(rows, 2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill:
                cell.fill = fill

    col_widths = [7, 11, 7, 16, 22, 14, 12, 14, 14, 13, 12,
                  16, 28, 14, 14, 12, 14, 16, 18]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_to_sheets_sync(rows: list, tab_name: str) -> str:
    """Sync function — runs in thread executor. Receives plain Python rows."""
    import gspread
    from google.oauth2.service_account import Credentials

    creds = Credentials.from_service_account_file(
        settings.GOOGLE_CREDENTIALS_PATH,
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    gc = gspread.authorize(creds)
    sh = gc.open_by_key(settings.GOOGLE_SPREADSHEET_ID)

    try:
        ws = sh.worksheet(tab_name)
        ws.clear()
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title=tab_name, rows=max(len(rows) + 5, 20), cols=len(HEADERS))

    ws.update([HEADERS], "A1")

    if rows:
        rows_str = [[str(v) if v != "" else "" for v in row] for row in rows]
        ws.update(rows_str, "A2")

    return sh.url


async def run_daily_export(bot: Bot) -> None:
    if not settings.EXPORT_CHAT_ID:
        return

    tz = pytz.timezone(settings.TIMEZONE)
    now = datetime.now(tz)
    date_from = now.replace(hour=0, minute=0, second=0, microsecond=0)
    date_to = now

    # Serialize ticket data while session is still open to avoid detached-instance errors
    async with async_session_factory() as session:
        tickets = await get_export_tickets(session, date_from, date_to)
        rows = [_ticket_to_row(t, tz) for t in tickets]

    period_str = date_from.strftime("%d.%m.%Y")

    if not rows:
        await bot.send_message(
            settings.EXPORT_CHAT_ID,
            f"📊 Выгрузка за {period_str}: нет обращений."
        )
        return

    tab_name = date_from.strftime("%d.%m.%Y")
    use_sheets = bool(settings.GOOGLE_CREDENTIALS_PATH and settings.GOOGLE_SPREADSHEET_ID)

    if use_sheets:
        try:
            loop = asyncio.get_event_loop()
            url = await loop.run_in_executor(
                None, _write_to_sheets_sync, rows, tab_name
            )
            await bot.send_message(
                settings.EXPORT_CHAT_ID,
                f"📊 Выгрузка за {period_str}\n"
                f"Обращений: {len(rows)}\n\n"
                f"📎 <a href=\"{url}\">Открыть Google Таблицу</a>",
                parse_mode="HTML",
            )
        except Exception as e:
            excel_bytes = _generate_excel(rows)
            filename = f"crm_export_{date_from.strftime('%d%m%Y')}.xlsx"
            await bot.send_document(
                settings.EXPORT_CHAT_ID,
                BufferedInputFile(excel_bytes, filename=filename),
                caption=f"📊 Выгрузка за {period_str}\nОбращений: {len(rows)}\n⚠️ Google Sheets недоступен: {e}",
            )
    else:
        excel_bytes = _generate_excel(rows)
        filename = f"crm_export_{date_from.strftime('%d%m%Y')}.xlsx"
        await bot.send_document(
            settings.EXPORT_CHAT_ID,
            BufferedInputFile(excel_bytes, filename=filename),
            caption=f"📊 Выгрузка за {period_str}\nОбращений: {len(rows)}",
        )
